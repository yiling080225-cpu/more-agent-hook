#!/usr/bin/env python3
"""采购清单生成脚本"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = Workbook()

# ============================================================
# Sheet 1: 采购清单
# ============================================================
ws = wb.active
ws.title = '采购清单'

hdr_font = Font(name='Arial', bold=True, size=11, color='FFFFFF')
hdr_fill = PatternFill('solid', fgColor='2F5496')
sub_fill = PatternFill('solid', fgColor='D6E4F0')
sub_font = Font(name='Arial', bold=True, size=10, color='2F5496')
normal_font = Font(name='Arial', size=10)
bold_font = Font(name='Arial', bold=True, size=10)
total_fill = PatternFill('solid', fgColor='FFF2CC')
total_font = Font(name='Arial', bold=True, size=11, color='C00000')
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left_wrap = Alignment(horizontal='left', vertical='center', wrap_text=True)

col_widths = {'A': 5, 'B': 24, 'C': 30, 'D': 6, 'E': 6, 'F': 10, 'G': 12, 'H': 24, 'I': 14}
for col, w in col_widths.items():
    ws.column_dimensions[col].width = w

ws.merge_cells('A1:I1')
ws['A1'] = '商用烟熏腊肉香肠炉 — 采购清单 (预算 30000元)'
ws['A1'].font = Font(name='Arial', bold=True, size=14, color='2F5496')
ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 30

headers = ['序号', '名称', '规格/型号', '单位', '数量', '单价(元)', '金额(元)', '推荐品牌/供应商', '备注']
for j, h in enumerate(headers, 1):
    cell = ws.cell(row=3, column=j, value=h)
    cell.font = hdr_font
    cell.fill = hdr_fill
    cell.alignment = center
    cell.border = thin_border
ws.row_dimensions[3].height = 22

items = [
    ('一、箱体结构材料', [
        ('304不锈钢板1.5mm', '1219x2438x1.5mm 2B食品级', '张', 4, 850, '太钢/宝钢/联众', '内胆+集油盘+发烟器'),
        ('201不锈钢板1.0mm', '1219x2438x1.0mm 2B表面', '张', 4, 450, '联众/宏旺', '外层箱体+门板'),
        ('岩棉保温板50mm', '1200x600x50mm 100kg/m3 300度', '块', 12, 45, '河北华能/神州', '6面保温层'),
        ('304不锈钢方管25x25x1.5', '6m/根', '根', 4, 85, '佛山市场', '箱体骨架+挂架横梁'),
        ('304不锈钢圆管D25x1.5', '6m/根', '根', 2, 75, '佛山市场', '挂架杆'),
        ('304不锈钢圆棒D12', '6m/根', '根', 1, 55, '佛山市场', '自制挂肉钩'),
        ('耐高温硅胶密封条', '15x10mm D型 250度自粘', '米', 15, 18, '河北橡塑/3M', '门框密封'),
        ('不锈钢门铰链 重型', '304 承重50kg', '个', 6, 25, '海蒂诗/DGN', '双开门'),
        ('不锈钢门把手300mm', '304大拉手', '个', 2, 45, '佛山市场', ''),
        ('钢化玻璃观察窗', '200x150x5mm 200度', '块', 2, 65, '信义/南玻', '门上观察窗'),
        ('不锈钢小型铰链', '304 50x50mm', '个', 2, 12, '海蒂诗', '发烟器盖'),
    ]),
    ('二、加热系统', [
        ('U型不锈钢电加热管', '3kW/380V 304护套 600mm', '支', 2, 280, '江苏兴泰', '箱体主加热'),
        ('电热板(发烟器)', '1kW/220V 150x150mm', '块', 1, 120, '江苏加热管厂', '木屑闷烧'),
        ('304冲孔网防护罩', '500x200mm 孔径5mm', '块', 2, 55, '安平丝网', '防滴油溅射'),
    ]),
    ('三、控制系统', [
        ('PID温控器', '宇电AI-518 精度0.1级 SSR输出', '台', 1, 580, '厦门宇电', '主控温,程序升温'),
        ('PT100温度传感器', 'D6x150mm 高温线3m', '支', 1, 85, '北京长城', '内腔测温'),
        ('湿度传感器', 'AM2305 高温型 0-100%', '支', 1, 95, '建大仁科', '内腔湿度'),
        ('机械式温控开关', '常闭型 150度断开', '个', 2, 45, '韩国彩虹', '超温保护串主回路'),
        ('24小时定时器', 'DH48S-S 导轨式', '个', 1, 75, '欣灵电气', '熏制时间设定'),
        ('漏电保护断路器', '正泰NXBLE-63 C63 30mA 4P', '个', 1, 125, '正泰', '主电源'),
        ('交流接触器', '正泰CJX2-2510 380V线圈', '个', 2, 65, '正泰', '加热管通断'),
        ('固态继电器SSR', 'SSR-40DA 40A DC-AC', '个', 2, 55, 'FOTEK', 'PID输出控制'),
        ('急停按钮', 'LAY39-11ZS 红色自锁1NC', '个', 1, 35, '正泰', '面板醒目位'),
        ('指示灯 LED 三色', 'AD16-22DS 220V 绿黄红', '个', 6, 8, '正泰', '各2个'),
        ('旋钮开关2档自锁', 'LAY39-11X', '个', 2, 15, '正泰', '风机/发烟独立'),
        ('304不锈控制箱', '300x200x150mm IP65', '个', 1, 180, '温州电箱', '控制面板壳'),
        ('耐高温硅胶电线', '2.5mm2 镀锡铜 200度', '米', 20, 8, '特软硅胶线', '加热管接线'),
        ('BV铜芯电线', '2.5mm2', '米', 30, 4, '正泰/远东', '控制回路'),
        ('接线端子排+导轨', 'TB-1510 15A', '套', 1, 45, '正泰', ''),
    ]),
    ('四、循环系统', [
        ('耐高温轴流风机', 'CY125 150度 60W 380V', '台', 1, 380, '全风/九洲普惠', '背部内循环'),
        ('304不锈导流板', '1.0mm折弯 600x300mm', '块', 1, 80, '钣金加工', '风机出口'),
    ]),
    ('五、发烟器系统', [
        ('304不锈发烟筒体', 'D180xH250mm 1.5mm定制', '个', 1, 350, '不锈加工', '发烟器主体'),
        ('DN50不锈钢烟管', '304 D50x1.5mm含接头', '米', 2, 45, '佛山市场', '发烟器到箱体'),
        ('DN50不锈钢球阀', '304内螺纹', '个', 1, 65, '温州阀门', '烟量调节'),
    ]),
    ('六、排烟及环保系统', [
        ('DN80不锈钢排烟管', '304 D80x1.5mm', '米', 3, 60, '佛山市场', '排烟主管'),
        ('DN80手动蝶阀', '304法兰式', '个', 1, 120, '温州阀门', '排烟调节'),
        ('水幕过滤箱体', '304 2.0mm 800x350x500mm', '个', 1, 850, '不锈加工定制', '三级水洗主体'),
        ('不锈钢喷淋球头', '304 DN15螺旋', '个', 6, 25, '温州喷淋', '每级2个'),
        ('循环水泵', 'MP-20R 50W 5m扬程耐酸碱', '台', 1, 220, '西山泵业', '水幕循环'),
        ('蜂窝活性炭', '100x100x100 防水碘值800', '块', 10, 15, '活性炭厂', '末端过滤层'),
        ('PP棉滤芯 10寸', '5um标准', '支', 2, 12, '通用', '水泵前过滤'),
        ('DN15不锈水管+接头套', '304含弯头三通对丝', '套', 1, 120, '温州管件', '水循环管路'),
        ('PE废水收集槽 20L', 'PE带盖', '个', 1, 55, '塑料市场', '焦油废水'),
    ]),
    ('七、结构附件', [
        ('万向轮带刹车6寸', '重型PU 单轮200kg', '只', 4, 85, '诺力/向荣', '底部安装'),
        ('304排污球阀1/2寸', '内螺纹', '个', 1, 32, '温州阀门', '底部排污'),
        ('不锈钢蝶形螺栓M8x25', '304', '套', 20, 3.5, '东明/固万基', '挂架固定'),
        ('不锈钢铆钉D4x10', '304', '盒', 1, 25, '固万基', '箱体组装'),
        ('耐高温密封胶', '硅酮300度 食品级', '支', 3, 28, '瓦克/道康宁', '接缝密封'),
    ]),
    ('八、加工费用', [
        ('激光切割+折弯加工', '含编程 6面+门+零件', '批', 1, 2500, '钣金加工厂', '主材加工'),
        ('氩弧焊接', '箱体+发烟+过滤箱', '批', 1, 2800, '不锈焊接', '食品级焊接'),
        ('电气组装调试', '含布线/编程/试机', '批', 1, 1500, '电控工程师', '持电工证'),
        ('物流运费+包装', '含木架 约300kg', '批', 1, 800, '专线物流', '省内'),
        ('预留不可预见费', '耗材/小件/返工', '批', 1, 1500, '', '5%预留'),
    ]),
]

row = 4
item_num = 1
section_totals = []
grand_total = 0

for section, sub_items in items:
    ws.merge_cells(f'A{row}:I{row}')
    ws.cell(row=row, column=1, value=section).font = sub_font
    ws.cell(row=row, column=1).fill = sub_fill
    ws.cell(row=row, column=1).alignment = left_wrap
    for j in range(1, 10):
        ws.cell(row=row, column=j).border = thin_border
        ws.cell(row=row, column=j).fill = sub_fill
    ws.row_dimensions[row].height = 22
    row += 1

    section_sum = 0
    for name, spec, unit, qty, price, brand, note in sub_items:
        amount = qty * price
        section_sum += amount
        vals = [item_num, name, spec, unit, qty, price, amount, brand, note]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=j, value=v)
            cell.font = normal_font
            cell.border = thin_border
            cell.alignment = center if j in (1, 4, 5, 6, 7) else left_wrap
            if j == 6:
                cell.number_format = '#,##0'
            if j == 7:
                cell.number_format = '#,##0'
        ws.row_dimensions[row].height = 20
        item_num += 1
        row += 1

    ws.merge_cells(f'A{row}:F{row}')
    ws.cell(row=row, column=1, value=f'{section} 小计').font = bold_font
    ws.cell(row=row, column=1).alignment = Alignment(horizontal='right', vertical='center')
    ws.cell(row=row, column=7, value=section_sum).font = bold_font
    ws.cell(row=row, column=7).number_format = '#,##0'
    ws.cell(row=row, column=7).alignment = center
    for j in range(1, 10):
        ws.cell(row=row, column=j).border = thin_border
    ws.row_dimensions[row].height = 22
    section_totals.append((section, section_sum))
    grand_total += section_sum
    row += 2

# Grand total
ws.merge_cells(f'A{row}:F{row}')
ws.cell(row=row, column=1, value='采购总计').font = total_font
ws.cell(row=row, column=1).fill = total_fill
ws.cell(row=row, column=1).alignment = Alignment(horizontal='right', vertical='center')
ws.cell(row=row, column=7, value=grand_total).font = total_font
ws.cell(row=row, column=7).fill = total_fill
ws.cell(row=row, column=7).number_format = '#,##0'
ws.cell(row=row, column=7).alignment = center
remain = 30000 - grand_total
ws.cell(row=row, column=9, value=f'预算30000 剩余{remain}').font = total_font
ws.cell(row=row, column=9).fill = total_fill
for j in range(1, 10):
    ws.cell(row=row, column=j).border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='medium'), bottom=Side(style='medium')
    )
ws.row_dimensions[row].height = 26

# ============================================================
# Sheet 2: 分类汇总
# ============================================================
ws2 = wb.create_sheet('分类汇总')
ws2.column_dimensions['A'].width = 30
ws2.column_dimensions['B'].width = 16
ws2.column_dimensions['C'].width = 16

ws2.merge_cells('A1:C1')
ws2['A1'] = '预算分类汇总'
ws2['A1'].font = Font(name='Arial', bold=True, size=14, color='2F5496')
ws2['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws2.row_dimensions[1].height = 28

for j, h in enumerate(['类别', '金额(元)', '占比'], 1):
    cell = ws2.cell(row=3, column=j, value=h)
    cell.font = hdr_font
    cell.fill = hdr_fill
    cell.alignment = center
    cell.border = thin_border

for i, (sec, amt) in enumerate(section_totals):
    r = 4 + i
    ws2.cell(row=r, column=1, value=sec).font = normal_font
    ws2.cell(row=r, column=1).alignment = left_wrap
    ws2.cell(row=r, column=2, value=amt).font = normal_font
    ws2.cell(row=r, column=2).number_format = '#,##0'
    ws2.cell(row=r, column=2).alignment = center
    ws2.cell(row=r, column=3).font = normal_font
    ws2.cell(row=r, column=3).number_format = '0.0%'
    ws2.cell(row=r, column=3).alignment = center
    tr2 = 4 + len(section_totals)
    ws2[f'C{r}'] = f'=B{r}/B{tr2}'
    for j in range(1, 4):
        ws2.cell(row=r, column=j).border = thin_border

tr = 4 + len(section_totals)
ws2.cell(row=tr, column=1, value='合计').font = total_font
ws2.cell(row=tr, column=1).fill = total_fill
ws2.cell(row=tr, column=1).alignment = Alignment(horizontal='right', vertical='center')
ws2.cell(row=tr, column=2, value=grand_total).font = total_font
ws2.cell(row=tr, column=2).fill = total_fill
ws2.cell(row=tr, column=2).number_format = '#,##0'
ws2.cell(row=tr, column=2).alignment = center
ws2.cell(row=tr, column=3).font = total_font
ws2.cell(row=tr, column=3).fill = total_fill
ws2.cell(row=tr, column=3).number_format = '0.0%'
ws2.cell(row=tr, column=3).alignment = center
for j in range(1, 4):
    ws2.cell(row=tr, column=j).border = thin_border

# ============================================================
# Sheet 3: 说明
# ============================================================
ws3 = wb.create_sheet('关键说明')
ws3.column_dimensions['A'].width = 5
ws3.column_dimensions['B'].width = 75

notes = [
    ('', ''),
    ('', '一、设计参数'),
    ('', '外形尺寸: 1200x800x1800mm 有效容积约1.0m3 批次容量50-80kg'),
    ('', '总功率: 7.06kW(380V) = 加热6kW + 发烟1kW + 风机0.06kW'),
    ('', '温度范围: 室温至120度 精度正负1度 PID程序控温'),
    ('', '保温: 50mm岩棉 密度100kg/m3 耐温300度'),
    ('', ''),
    ('', '二、环保设计'),
    ('', '1. 水幕过滤箱: 三级水洗降尘 + 活性炭吸附 + PP棉精滤'),
    ('', '2. 烟气净化: PM2.5去除率大于等于85% 焦油去除率大于等于80%'),
    ('', '3. 排放标准: 符合GB 18483-2001饮食业油烟排放标准'),
    ('', '4. 废水收集: 焦油废水独立收集 定期交危废处理单位清运'),
    ('', '5. 能源效率: 50mm岩棉保温 热损失减少大于等于80%'),
    ('', '6. 木屑优化: 电热闷烧比直接燃烧减少木屑浪费大于等于30%'),
    ('', '7. 噪声控制: 风机噪声小于等于55dB(A)'),
    ('', ''),
    ('', '三、安全设计'),
    ('', '1. 独立机械温控150度超温断电(串入主回路)'),
    ('', '2. 漏电保护30mA动作电流 4P断路器'),
    ('', '3. 急停按钮面板正面 电气防护等级大于等于IP54'),
    ('', '4. 加热管接线使用耐高温硅胶线(200度)'),
    ('', '5. 箱体可靠接地'),
    ('', ''),
    ('', '四、维护周期'),
    ('', '每批: 清理集油盘 | 每周: 查水幕水位水质 | 每月: 查密封条'),
    ('', '每季: 清烟道积焦 | 半年: 校探头+换活性炭 | 每年: 风机加高温脂'),
    ('', ''),
    ('', '五、采购渠道建议'),
    ('', '不锈钢板材: 佛山集中采购价格最优(太钢/宝钢)'),
    ('', '电气元件: 1688/淘宝工业品 正泰性价比高 欧姆龙更可靠'),
    ('', '定制钣金: 找有食品设备加工经验的厂家'),
    ('', '焊接: 需304不锈钢氩弧焊经验 焊缝需酸洗钝化'),
    ('', ''),
    ('', '六、工期估算(约3-4周)'),
    ('', '板材采购+切割: 5-7天 | 其他采购: 3-5天(并行)'),
    ('', '钣金折弯+焊接: 5-7天 | 电气组装+调试: 3-5天'),
]

for i, (col_a, col_b) in enumerate(notes):
    r = i + 1
    ws3.cell(row=r, column=1, value=col_a)
    ws3.cell(row=r, column=2, value=col_b)
    if col_b and any(col_b.startswith(p) for p in ['一、', '二、', '三、', '四、', '五、', '六、']):
        ws3.cell(row=r, column=2).font = Font(name='Arial', bold=True, size=12, color='2F5496')
    elif col_b:
        ws3.cell(row=r, column=2).font = Font(name='Arial', size=10)

ws.freeze_panes = 'A4'

output_path = r'd:\multimodal_agent_federation_mvp\preview\smokehouse_procurement.xlsx'
wb.save(output_path)
print(f'[OK] {output_path}')
print(f'Items: {item_num-1}')
for sec, amt in section_totals:
    print(f'  {sec}: {amt:,}')
print(f'  TOTAL: {grand_total:,} / 30000 = remain {30000-grand_total:,}')
